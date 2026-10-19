#!/usr/bin/env python3
"""Independent checker (the R4 oracle) for the XLSX export, written from the specification alone.

Spec: PLAN.md section 5 "Export" (the layout contract) and docs/spec/REQUIREMENTS.md REQ-EXP-001 to 010. The workbook
is read by its English labels; every Summary figure is recomputed from the Plan vs done rows and compared.

Checks (every difference is printed with its sheet, row and column)
    1. sheets `Plan vs done` and `Summary`, in this order; each starts with the header block (labels in column A,
       values in B, then one empty row); both blocks hold the same values; Planned sessions and Planned hours equal
       the count and the hours of the planned rows
    2. Plan vs done: the 14 column labels; dates as yyyy-mm-dd text in date order, unplanned sessions after the
       planned days of their date; one of the six status texts; planned hours a number on a planned day and empty on
       an `Added` row; days late a whole number >= 0 exactly on the rows with an entry and equal to Written on minus
       Date (REQ-LOG-004); `Not logged` and `Upcoming` rows hold no entry field; every `Not logged` day before every
       `Upcoming` day; with --today, no `Upcoming` on or before it and no `Not logged` after it
    3. Summary: the three tables with their titles and column labels; every cell of every group row recomputed from
       the rows; groups missing, repeated or out of order; each table's rows add up to Overall in every count and
       hours column; planned = done + partly done + not done + not logged + upcoming on every row
    4. no cell anywhere holds a formula, a link, a date or a boolean: every cell is empty, a number or a string

Assumptions (literal readings where the spec leaves room)
    - A row is an unplanned session exactly when its status is `Added`; every other row is a planned day.
    - Days late counts calendar days from Date to Written on, 0 when written on or before the date (REQ-LOG-004:
      "late by N days"); Written on is yyyy-mm-dd text like every date in the file.
    - Planned hours on a planned day is a number >= 0 with at most two decimals, or empty (a day without hours,
      warning missing_hours), which counts as 0.
    - Continues in the next session holds `yes` (any letter case) or nothing.
    - Per block, `Added` rows count in `Unplanned sessions` whatever their Block cell, and a planned day with an
      empty Block in `No block`. Per evaluation period, `Added` rows count in their period like planned days.
    - Every Per block row and `Outside evaluation periods` are listed only when some row belongs to them (blocks are
      listed by their first appearance in the rows, so a block without rows has none); a named evaluation period
      may be listed with no rows and then shows zeros with an empty coverage.
    - Blocks are listed in order of first appearance in the Plan vs done rows. Period dates are not in the file, so
      periods with rows are checked to be listed in the order of their first row; a period with no rows can sit
      anywhere among them.
    - Coverage is rounded half up on the exact fraction. Hours are exact (written with up to three decimals, the
      stored hours having two): every hours cell, and every "adds up" sum, must equal the exact value to 1e-9.
    - Most days late is a maximum, not a count, so it is recomputed per group but not added up.
    - The header values Planned sessions and Planned hours may be numbers or numeric text; every header value must
      be non-empty and identical on both sheets (they come from one export).
    - "Today" is not in the file: without --today only the order of `Not logged` and `Upcoming` days is checked.

Usage
    python3 tools/check_export.py export.xlsx [--today 2026-10-19]
    python3 tools/check_export.py --selftest       # a correct synthetic workbook passes, every mutant fails
"""
import argparse
import copy
import datetime as dt
import math
import re
import sys
import tempfile
from dataclasses import dataclass
from decimal import Decimal
from fractions import Fraction
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROWS, SUMMARY = "Plan vs done", "Summary"
HEADER = ["Module", "Course", "Activity", "School year", "Plan dates", "Planned sessions", "Planned hours",
          "Exported by", "Exported on"]
COLUMNS = ["Date", "Evaluation period", "Block", "Planned", "Planned hours", "Status", "What was done", "Reason",
           "Reason note", "Corrective measure", "Continues in the next session", "Written on", "Days late", "Logged by"]
ENTRY_FIELDS = COLUMNS[6:]
DONE, PARTLY, NOT_DONE, ADDED, NOT_LOGGED, UPCOMING = "Done", "Partly done", "Not done", "Added", "Not logged", "Upcoming"
STATUSES = (DONE, PARTLY, NOT_DONE, ADDED, NOT_LOGGED, UPCOMING)
OVERALL, PER_BLOCK, PER_PERIOD = TABLES = ("Overall", "Per block", "Per evaluation period")
FIGURES = ["Group", "Planned", "Done", "Partly done", "Not done", "Not logged", "Upcoming", "Added", "Coverage (%)",
           "Planned hours", "Hours done", "Hours not done", "Written late", "Most days late"]
STATES = FIGURES[2:7]
ADDS_UP = FIGURES[1:8] + ["Planned hours", "Hours done", "Hours not done", "Written late"]
HOURS = ("Planned hours", "Hours done", "Hours not done")
ALL, NO_BLOCK, UNPLANNED, OUTSIDE = "All sessions", "No block", "Unplanned sessions", "Outside evaluation periods"
GROUP_OF = {OVERALL: lambda x: ALL,
            PER_BLOCK: lambda x: UNPLANNED if x.status == ADDED else x.block or NO_BLOCK,
            PER_PERIOD: lambda x: x.period or OUTSIDE}
LISTED_LAST = {OVERALL: (), PER_BLOCK: (NO_BLOCK, UNPLANNED), PER_PERIOD: (OUTSIDE,)}
ISO_DATE = re.compile(r"^[0-9]{4}-[0-9]{2}-[0-9]{2}$")
CENT = Decimal("0.01")
EXACT = Decimal("1e-9")  # float noise only: the hours are exact decimals


@dataclass
class Row:
    r: int
    date: object
    status: object
    period: object
    block: object
    hours: Decimal
    late: object


class Sheet:
    """The values of a worksheet, read once, addressed 1-based like Excel."""

    def __init__(self, ws):
        self.title = ws.title
        width = max(ws.max_column, len(COLUMNS))
        self.rows = [list(r) + [None] * (width - len(r)) for r in ws.iter_rows(values_only=True)]

    def get(self, r, c):
        return self.rows[r - 1][c - 1] if 1 <= r <= len(self.rows) else None

    def empty_row(self, r):
        return not 1 <= r <= len(self.rows) or all(empty(v) for v in self.rows[r - 1])

    def find(self, text):
        return next((r for r in range(1, len(self.rows) + 1) if self.get(r, 1) == text), None)

    def at(self, r, c):
        return f"{self.title}!{get_column_letter(c)}{r}"


def empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def dec(v):
    return Decimal(repr(v)) if isinstance(v, float) else Decimal(v)


def to_date(v):
    if isinstance(v, str) and ISO_DATE.match(v):
        try:
            return dt.date.fromisoformat(v)
        except ValueError:
            return None
    return None


def number(v):
    if is_number(v):
        return dec(v)
    try:
        return Decimal(v.strip().replace(",", "."))
    except Exception:  # not numeric text
        return None


def matches(got, want, hours):
    if want is None:
        return empty(got)
    if not is_number(got):
        return False
    return abs(dec(got) - want) <= EXACT if hours else got == want


def figures(xs):
    """The Summary figures of REQ-EXP-005 for a set of rows."""
    n = {s: sum(1 for x in xs if x.status == s) for s in STATUSES}
    h = {s: sum((x.hours for x in xs if x.status == s), Decimal(0)) for s in STATUSES}
    planned = [x for x in xs if x.status != ADDED]
    divisor = n[DONE] + n[PARTLY] + n[NOT_DONE] + n[NOT_LOGGED]
    late = [x.late for x in xs if x.late]
    return {"Planned": len(planned), "Done": n[DONE], "Partly done": n[PARTLY], "Not done": n[NOT_DONE],
            "Not logged": n[NOT_LOGGED], "Upcoming": n[UPCOMING], "Added": n[ADDED],
            "Coverage (%)": math.floor(Fraction(100 * (2 * n[DONE] + n[PARTLY]), 2 * divisor) + Fraction(1, 2))
            if divisor else None,
            "Planned hours": sum((x.hours for x in planned), Decimal(0)),
            "Hours done": h[DONE] + h[PARTLY] / 2, "Hours not done": h[NOT_DONE] + h[PARTLY] / 2,
            "Written late": len(late), "Most days late": max(late, default=0)}


# ── reading and checking ────────────────────────────────────────────────────────────────────────────────────────
def check_cells(wb, out):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if c.data_type == "f" or (not isinstance(v, str) and str(v).startswith("=")):
                    out.append(f"{ws.title}!{c.coordinate}: holds a formula {str(v)!r}; figures are written as numbers")
                elif v is not None and not is_number(v) and not isinstance(v, str):
                    out.append(f"{ws.title}!{c.coordinate}: holds a {type(v).__name__} ({v!r}), not a number or a string")
                if getattr(c, "hyperlink", None):
                    out.append(f"{ws.title}!{c.coordinate}: is a link; addresses stay plain text")


def read_header(sh, stop, out):
    """{label: (row, value)} of the header block, the rows above row `stop`."""
    if stop is None:  # no table found: the block ends at the first empty row
        stop = next(r for r in range(1, len(sh.rows) + 2) if sh.empty_row(r)) + 1
    found = {}
    for r in range(1, stop):
        label = sh.get(r, 1)
        if empty(label):
            continue
        if label in HEADER and label not in found:
            found[label] = (r, sh.get(r, 2))
        else:
            out.append(f"{sh.at(r, 1)}: unexpected label {label!r} in the header block")
    for label in HEADER:
        if label not in found:
            out.append(f"{sh.title}: missing header label {label!r} in column A")
        elif empty(found[label][1]):
            out.append(f"{sh.at(found[label][0], 2)}: {label} has no value")
    order = [found[label][0] for label in HEADER if label in found]
    if order != sorted(order):
        out.append(f"{sh.title}: header labels are not in the order {', '.join(HEADER)}")
    if stop > 1 and not (sh.empty_row(stop - 1) and not sh.empty_row(stop - 2)):
        out.append(f"{sh.at(stop - 1, 1)}: expected exactly one empty row between the header block and the table")
    return found


def read_rows(sh, start, today, out):
    """The Plan vs done rows, checked one by one; None when the table cannot be read."""
    if start is None:
        out.append(f"{ROWS}: no row-table header (a row with 'Date' in column A)")
        return None
    bad = [c for c, want in enumerate(COLUMNS, 1) if sh.get(start, c) != want]
    for c in bad:
        out.append(f"{sh.at(start, c)}: header label expected {COLUMNS[c - 1]!r}, got {sh.get(start, c)!r}")
    if bad:
        out.append(f"{ROWS}: rows not checked, their header row does not match the contract")
        return None
    rows, prev = [], None
    for r in range(start + 1, len(sh.rows) + 1):
        if sh.empty_row(r):
            continue
        v = dict(zip(COLUMNS, sh.rows[r - 1]))

        def at(col):
            return sh.at(r, COLUMNS.index(col) + 1)

        d, status, hours, late = to_date(v["Date"]), v["Status"], v["Planned hours"], v["Days late"]
        added, entry = status == ADDED, status in (DONE, PARTLY, NOT_DONE, ADDED)
        good_hours = is_number(hours) and hours >= 0 and dec(hours) == dec(hours).quantize(CENT)
        if d is None:
            out.append(f"{at('Date')}: date {v['Date']!r} is not yyyy-mm-dd text")
        if status not in STATUSES:
            out.append(f"{at('Status')}: unknown status {status!r}, expected one of {', '.join(STATUSES)}")
        if added and not empty(hours):
            out.append(f"{at('Planned hours')}: an unplanned session ('Added') has planned hours {hours!r}")
        elif not added and not empty(hours) and not good_hours:
            out.append(f"{at('Planned hours')}: a planned day has hours >= 0 with at most two decimals, or none, "
                       f"got {hours!r}")
        if entry:
            if not (is_number(late) and late >= 0 and late == int(late)):
                out.append(f"{at('Days late')}: a row with an entry needs a whole number of days late >= 0, got {late!r}")
                late = None
            written = to_date(v["Written on"])
            if written is None:
                out.append(f"{at('Written on')}: {v['Written on']!r} is not yyyy-mm-dd text")
            elif d and late is not None and late != max(0, (written - d).days):
                out.append(f"{at('Days late')}: Days late is {late!r}, but Written on {written} is "
                           f"{max(0, (written - d).days)} day(s) late for {d}")
            if empty(v["Logged by"]):
                out.append(f"{at('Logged by')}: a row with an entry names who logged it")
            cont = v["Continues in the next session"]
            if not empty(cont) and str(cont).strip().lower() != "yes":
                out.append(f"{at('Continues in the next session')}: expected 'yes' or empty, got {cont!r}")
        elif status in (NOT_LOGGED, UPCOMING):
            for f in ENTRY_FIELDS:
                if not empty(v[f]):
                    out.append(f"{at(f)}: the row is {status!r} (no entry) but {f} holds {v[f]!r}")
            late = None
        if today and d:
            if status == UPCOMING and d <= today:
                out.append(f"{at('Status')}: 'Upcoming' on {d}, which is not after today ({today})")
            if status == NOT_LOGGED and d > today:
                out.append(f"{at('Status')}: 'Not logged' on {d}, which is after today ({today})")
        if d and prev:
            if d < prev[0]:
                out.append(f"{at('Date')}: {d} comes after {prev[0]}; rows are not in date order")
            elif d == prev[0] and prev[1] and not added:
                out.append(f"{at('Date')}: a planned day after an unplanned session of the same date")
        if d:
            prev = (d, added)
        rows.append(Row(r, d, status, None if empty(v["Evaluation period"]) else v["Evaluation period"],
                        None if empty(v["Block"]) else v["Block"],
                        dec(hours) if not added and is_number(hours) else Decimal(0),
                        int(late) if is_number(late) else None))
    not_logged = [x.date for x in rows if x.status == NOT_LOGGED and x.date]
    upcoming = [x.date for x in rows if x.status == UPCOMING and x.date]
    if not_logged and upcoming and max(not_logged) >= min(upcoming):
        out.append(f"{ROWS}: 'Not logged' on {max(not_logged)} is not before 'Upcoming' on {min(upcoming)}; "
                   f"one today cannot give both")
    return rows


def read_tables(sh, titles, out):
    """{title: [(row, {figure: value})]} of the Summary tables whose header row matches."""
    tables = {}
    for t, r in zip(TABLES, titles):
        if r is None:
            out.append(f"{SUMMARY}: missing table title {t!r} in column A")
            continue
        bad = [c for c, want in enumerate(FIGURES, 1) if sh.get(r + 1, c) != want]
        for c in bad:
            out.append(f"{sh.at(r + 1, c)}: header label expected {FIGURES[c - 1]!r}, got {sh.get(r + 1, c)!r}")
        if bad:
            continue
        rr, groups = r + 2, []
        while not sh.empty_row(rr) and sh.get(rr, 1) not in TABLES:
            groups.append((rr, dict(zip(FIGURES, sh.rows[rr - 1]))))
            rr += 1
        tables[t] = groups
    found = [r for r in titles if r]
    if found != sorted(found):
        out.append(f"{SUMMARY}: tables are not in the order {', '.join(TABLES)}")
    for t, r in zip(TABLES[1:], titles[1:]):
        if r and not (sh.empty_row(r - 1) and not sh.empty_row(r - 2)):
            out.append(f"{sh.at(r, 1)}: expected exactly one empty row before the title {t!r}")
    return tables


def check_summary(sh, tables, rows, out):
    overall = tables[OVERALL][0][1] if len(tables.get(OVERALL, [])) == 1 else None
    for t, groups in tables.items():
        members = {}  # insertion order = order of first appearance in the (date-ordered) rows
        for x in rows:
            members.setdefault(GROUP_OF[t](x), []).append(x)
        names = [g["Group"] for _, g in groups]
        if t == OVERALL and names != [ALL]:
            out.append(f"{SUMMARY}: table {OVERALL!r} must hold exactly one row {ALL!r}, found {names}")
        for name in sorted({str(n) for n in names if names.count(n) > 1}):
            out.append(f"{SUMMARY}: table {t!r} lists group {name!r} more than once")
        for name, xs in members.items():
            if name not in names:
                out.append(f"{SUMMARY}: table {t!r} has no row for group {name!r}, which has {len(xs)} row(s) in {ROWS}")
        for r, g in groups:
            name = g["Group"]
            xs = members.get(name, [])
            if not xs and (t == PER_BLOCK or name == OUTSIDE):
                out.append(f"{sh.at(r, 1)}: group {name!r} is listed but no row of {ROWS} belongs to it")
            want = figures(xs)
            for c, f in enumerate(FIGURES[1:], 2):
                if not matches(g[f], want[f], f in HOURS):
                    why = f" (no row of {ROWS} belongs to this group: every figure 0, coverage empty)" if not xs else ""
                    out.append(f"{sh.at(r, c)} ({t} / {name}, {f}): got {g[f]!r}, "
                               f"expected {'empty' if want[f] is None else want[f]}{why}")
            states = [g[f] for f in STATES]
            if all(is_number(v) for v in states + [g["Planned"]]) and g["Planned"] != sum(states):
                out.append(f"{sh.at(r, 2)} ({t} / {name}): Planned {g['Planned']!r} is not done + partly done + "
                           f"not done + not logged + upcoming = {sum(states)!r}")
        last = [n for n in LISTED_LAST[t] if n in names]
        named = [n for n in names if n not in LISTED_LAST[t]]
        if names != named + last:
            out.append(f"{SUMMARY}: table {t!r} lists {names}; {' then '.join(LISTED_LAST[t])} come last")
        first_seen = [n for n in members if n in named]
        if [n for n in named if n in members] != first_seen:
            out.append(f"{SUMMARY}: table {t!r} lists {[n for n in named if n in members]}, but they first appear in "
                       f"{ROWS} in the order {first_seen}")
        if t != OVERALL and overall:
            for f in ADDS_UP:
                vals = [g[f] for _, g in groups]
                if all(is_number(v) for v in vals + [overall[f]]):
                    total = sum((dec(v) for v in vals), Decimal(0))
                    if abs(total - dec(overall[f])) > EXACT:
                        out.append(f"{SUMMARY}: rows of {t!r} do not add up to Overall in column {f!r}: "
                                   f"{total} vs {overall[f]!r}")


def check(path, today=None):
    """Every difference between the workbook and the contract, as printable lines."""
    try:
        wb = openpyxl.load_workbook(path)  # not read_only, not data_only: formulas keep data_type 'f'
    except Exception as e:  # any reason the file is not a workbook
        return [f"{path}: not a readable XLSX workbook ({e})"]
    out = []
    if wb.sheetnames != [ROWS, SUMMARY]:
        out.append(f"workbook: sheets are {wb.sheetnames}, expected {[ROWS, SUMMARY]} in this order")
    check_cells(wb, out)
    if ROWS not in wb.sheetnames or SUMMARY not in wb.sheetnames:
        return out
    pvd, summary = Sheet(wb[ROWS]), Sheet(wb[SUMMARY])
    start = pvd.find("Date")
    titles = [summary.find(t) for t in TABLES]
    heads = {pvd.title: (pvd, read_header(pvd, start, out)),
             summary.title: (summary, read_header(summary, min([r for r in titles if r], default=None), out))}
    for label in HEADER:
        a, b = heads[ROWS][1].get(label), heads[SUMMARY][1].get(label)
        if a and b and a[1] != b[1]:
            out.append(f"{summary.at(b[0], 2)}: {label} is {b[1]!r}, but {a[1]!r} on {ROWS}")
    rows = read_rows(pvd, start, today, out)
    tables = read_tables(summary, titles, out)
    if rows is None:
        return out
    planned = [x for x in rows if x.status != ADDED]
    want = {"Planned sessions": Decimal(len(planned)),
            "Planned hours": sum((x.hours for x in planned), Decimal(0))}
    for sh, head in heads.values():
        for label, w in want.items():
            if label in head:
                r, v = head[label]
                got = number(v)
                if got is None or abs(got - w) > EXACT:
                    out.append(f"{sh.at(r, 2)}: {label} is {v!r}, but the planned rows of {ROWS} give {w}")
    check_summary(summary, tables, rows, out)
    return out


# ── self test ───────────────────────────────────────────────────────────────────────────────────────────────────
P1, P2 = "1a avaluació", "2a avaluació"
RA1, RA2, RA3 = "RA1 · Estructura", "RA2 · Codi", "RA3 · Proves"


def day(date, period, block, planned, hours, status, written=None, late=None, reason=None, cont=None):
    entry = written is not None
    return [date, period, block, planned, hours, status, "Sessió feta" if entry and status != NOT_DONE else None,
            reason, "Nota del motiu" if reason else None, "Es recupera" if status == NOT_DONE else None, cont,
            written, late, "Anna Puig" if entry else None]


# Today 2026-10-19. Outside every period: 09-11. RA2: 12.5 % coverage (13 half up, 12 half to even). Overall: 30 %,
# 43 % if not logged were left out of the divisor. Partly done 1.25 h gives 0.625 h: hours are written exactly.
FIXTURE = {
    "header": [["Module", "MP03 · Programació"], ["Course", "1r DAM"], ["Activity", "Seguiment MP03"],
               ["School year", "2026-09-07 to 2027-06-21"], ["Plan dates", "2026-09-11 to 2026-11-02"],
               ["Planned sessions", 12], ["Planned hours", 18], ["Exported by", "Anna Puig"],
               ["Exported on", "2026-10-19 10:30"]],
    "columns": COLUMNS,
    "rows": [
        day("2026-09-11", None, RA1, "Presentació\n• Objectius\n• Avaluació", 1, DONE, "2026-09-11", 0),
        day("2026-09-11", None, None, None, None, ADDED, "2026-09-12", 1),
        day("2026-09-14", P1, RA1, "Variables", 2, DONE, "2026-09-16", 2),
        day("2026-09-16", P1, RA1, "Tipus", 1.5, PARTLY, "2026-09-16", 0, "Falta de temps", "yes"),
        day("2026-09-16", P1, None, None, None, ADDED, "2026-09-16", 0),
        day("2026-09-21", P1, None, "Repàs", 0.75, NOT_DONE, "2026-09-21", 0, "Excursió"),
        day("2026-09-23", P1, RA2, "Bucles", 2, NOT_LOGGED),
        day("2026-09-30", P1, RA2, "Condicionals", 1, NOT_DONE, "2026-10-01", 1, "Malaltia"),
        day("2026-10-05", P1, RA2, "Funcions", 1.25, PARTLY, "2026-10-09", 4, "Falta de temps"),
        day("2026-10-14", P1, None, "Pràctica", 1, NOT_LOGGED),
        day("2026-10-19", P1, RA2, "Arrays", 2, NOT_LOGGED),
        day("2026-10-23", P1, RA2, "Fitxers", 2, UPCOMING),
        day("2026-10-26", P1, RA3, "Tests", 1.5, NOT_DONE, "2026-10-15", 0, "Vaga"),  # cancelled in advance
        day("2026-11-02", P1, RA3, "Depuració", 2, UPCOMING),
    ],
    "summary": [  # computed by hand, not by figures()
        (OVERALL, [[ALL, 12, 2, 2, 3, 3, 2, 2, 30, 18, 4.375, 4.625, 4, 4]]),
        (PER_BLOCK, [[RA1, 3, 2, 1, 0, 0, 0, 0, 83, 4.5, 3.75, 0.75, 1, 2],
                     [RA2, 5, 0, 1, 1, 2, 1, 0, 13, 8.25, 0.625, 1.625, 2, 4],
                     [RA3, 2, 0, 0, 1, 0, 1, 0, 0, 3.5, 0, 1.5, 0, 0],
                     [NO_BLOCK, 2, 0, 0, 1, 1, 0, 0, 0, 1.75, 0, 0.75, 0, 0],
                     [UNPLANNED, 0, 0, 0, 0, 0, 0, 2, None, 0, 0, 0, 1, 1]]),
        (PER_PERIOD, [[P1, 11, 1, 2, 3, 3, 2, 1, 22, 17, 3.375, 4.625, 3, 4],
                      [P2, 0, 0, 0, 0, 0, 0, 0, None, 0, 0, 0, 0, 0],
                      [OUTSIDE, 1, 1, 0, 0, 0, 0, 1, 100, 1, 1, 0, 1, 1]]),
    ],
}


def write_fixture(path, fx):
    wb = openpyxl.Workbook()
    pvd = wb.active
    pvd.title = ROWS
    summary = wb.create_sheet(SUMMARY)
    for ws in (pvd, summary):
        for row in fx["header"]:
            ws.append(row)
        ws.append([])
    pvd.append(fx["columns"])
    for row in fx["rows"]:
        pvd.append(row)
    for i, (title, groups) in enumerate(fx["summary"]):
        if i:
            summary.append([])
        summary.append([title])
        summary.append(FIGURES)
        for g in groups:
            summary.append(g)
    wb.save(path)


def put(fx, table, group, figure, value):
    for title, groups in fx["summary"]:
        for g in groups:
            if (title, g[0]) == (table, group):
                g[FIGURES.index(figure)] = value
    return fx


def drop(fx, table, group):
    for title, groups in fx["summary"]:
        if title == table:
            groups[:] = [g for g in groups if g[0] != group]
    return fx


def add(fx, table, before, group):
    for title, groups in fx["summary"]:
        if title == table:
            groups.insert(next(i for i, g in enumerate(groups) if g[0] == before), group)
    return fx


def cell(fx, i, column, value):
    fx["rows"][i][COLUMNS.index(column)] = value
    return fx


def swap(fx, i, j):
    fx["rows"][i], fx["rows"][j] = fx["rows"][j], fx["rows"][i]
    return fx


def relabel(fx, part, i, label):
    if part == "header":
        fx["header"][i][0] = label
    else:
        fx["columns"] = fx["columns"][:i] + [label] + fx["columns"][i + 1:]
    return fx


MUTANTS = [  # name, change to the correct fixture, a message that must appear, --today
    ("a wrong count", lambda f: put(f, OVERALL, ALL, "Done", 3), "(Overall / All sessions, Done): got 3, expected 2", None),
    ("coverage 12.5 rounded half to even", lambda f: put(f, PER_BLOCK, RA2, "Coverage (%)", 12),
     f"(Per block / {RA2}, Coverage (%)): got 12, expected 13", None),
    ("coverage divisor without not logged", lambda f: put(f, OVERALL, ALL, "Coverage (%)", 43),
     "(Overall / All sessions, Coverage (%)): got 43, expected 30", None),
    ("a formula cell", lambda f: put(f, OVERALL, ALL, "Planned", "=SUM(B17:B21)"), "holds a formula '=SUM(B17:B21)'", None),
    ("a missing header-block label", lambda f: relabel(f, "header", 6, "Hours"),
     "missing header label 'Planned hours'", None),
    ("a wrong column label", lambda f: relabel(f, "columns", 12, "Late days"),
     "header label expected 'Days late', got 'Late days'", None),
    ("per-block rows not adding up", lambda f: put(put(f, PER_BLOCK, RA3, "Not done", 2), PER_BLOCK, RA3, "Planned", 3),
     "rows of 'Per block' do not add up to Overall in column 'Not done'", None),
    ("hours off by 0.001", lambda f: put(f, PER_PERIOD, P1, "Hours done", 3.374),
     f"(Per evaluation period / {P1}, Hours done): got 3.374, expected 3.375", None),
    ("hours rounded to two decimals", lambda f: put(f, OVERALL, ALL, "Hours done", 4.38),
     "got 4.38, expected 4.375", None),
    ("planned hours with three decimals", lambda f: cell(f, 3, "Planned hours", 1.505),
     "a planned day has hours >= 0 with at most two decimals, or none, got 1.505", None),
    ("an empty period that is not zeros", lambda f: put(put(f, PER_PERIOD, P2, "Planned", 1), PER_PERIOD, P2, "Upcoming", 1),
     "Planned): got 1, expected 0 (no row of Plan vs done belongs to this group", None),
    ("a missing group row", lambda f: drop(f, PER_BLOCK, NO_BLOCK), "has no row for group 'No block'", None),
    ("a block listed with no rows", lambda f: add(f, PER_BLOCK, NO_BLOCK, ["RA4 · Extra"] + [0] * 7 + [None] + [0] * 5),
     "group 'RA4 · Extra' is listed but no row of Plan vs done belongs to it", None),
    ("No block listed with no rows", lambda f: cell(cell(f, 5, "Block", RA1), 9, "Block", RA1),
     "group 'No block' is listed but no row of Plan vs done belongs to it", None),
    ("Outside evaluation periods listed with no rows", lambda f: cell(cell(f, 0, "Evaluation period", P1), 1,
                                                                     "Evaluation period", P1),
     "group 'Outside evaluation periods' is listed but no row of Plan vs done belongs to it", None),
    ("an unknown status", lambda f: cell(f, 3, "Status", "Partly Done"), "unknown status 'Partly Done'", None),
    ("an Upcoming row with an entry field", lambda f: cell(f, 11, "What was done", "Avançat"),
     "the row is 'Upcoming' (no entry) but What was done holds 'Avançat'", None),
    ("rows out of date order", lambda f: swap(f, 2, 3), "comes after 2026-09-16; rows are not in date order", None),
    ("an entry without days late", lambda f: cell(f, 2, "Days late", None),
     "a row with an entry needs a whole number of days late >= 0, got None", None),
    ("days late not matching Written on", lambda f: cell(f, 8, "Days late", 3),
     "Days late is 3, but Written on 2026-10-09 is 4 day(s) late", None),
    ("an Added row with planned hours", lambda f: cell(f, 4, "Planned hours", 1),
     "an unplanned session ('Added') has planned hours 1", None),
    ("header Planned sessions off by one", lambda f: put_header(f, "Planned sessions", 13),
     "Planned sessions is 13, but the planned rows of Plan vs done give 12", None),
    ("a date written as a date cell", lambda f: cell(f, 0, "Date", dt.datetime(2026, 9, 11)),
     "holds a datetime", None),
    ("Upcoming on today", lambda f: f, "'Upcoming' on 2026-10-23, which is not after today", dt.date(2026, 10, 23)),
    ("Not logged after today", lambda f: f, "'Not logged' on 2026-10-19, which is after today", dt.date(2026, 10, 18)),
]


def put_header(fx, label, value):
    next(row for row in fx["header"] if row[0] == label)[1] = value
    return fx


def selftest():
    failures = 0
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "export.xlsx"

        def run(name, fx, today, want):
            nonlocal failures
            write_fixture(path, fx)
            got = check(path, today)
            hit = [m for m in got if want in m] if want else []
            ok = hit if want else not got
            print(("ok   " if ok else "FAIL ") + name + (f"\n       -> {hit[0]}" if hit else ""))
            if not ok:
                failures += 1
                print(f"       expected {repr(want) if want else 'no difference'}, got:")
                for m in got or ["(no difference)"]:
                    print(f"       {m}")

        run("correct workbook", copy.deepcopy(FIXTURE), None, None)
        run("correct workbook, --today 2026-10-19", copy.deepcopy(FIXTURE), dt.date(2026, 10, 19), None)
        run("correct workbook, --today 2026-10-22 (a Not done after today)", copy.deepcopy(FIXTURE),
            dt.date(2026, 10, 22), None)
        no_hours = cell(copy.deepcopy(FIXTURE), 9, "Planned hours", None)  # missing_hours: counts as 0
        for table, group, hours in ((OVERALL, ALL, 17), (PER_BLOCK, NO_BLOCK, 0.75), (PER_PERIOD, P1, 16)):
            put(no_hours, table, group, "Planned hours", hours)
        run("correct workbook, a planned day without hours", put_header(no_hours, "Planned hours", 17), None, None)
        for name, mutate, want, today in MUTANTS:
            run("mutant: " + name, mutate(copy.deepcopy(FIXTURE)), today, want)
    print("\n" + ("all passed" if not failures else f"{failures} failure(s)"))
    return 1 if failures else 0


def main():
    ap = argparse.ArgumentParser(description="Check an export XLSX against the layout contract (PLAN.md section 5).")
    ap.add_argument("file", nargs="?")
    ap.add_argument("--today", type=dt.date.fromisoformat, help="the exporting user's date, yyyy-mm-dd")
    ap.add_argument("--selftest", action="store_true")
    args = ap.parse_args()
    if args.selftest:
        sys.exit(selftest())
    if not args.file:
        ap.error("give an XLSX file or --selftest")
    problems = check(args.file, args.today)
    for p in problems:
        print(p)
    print(f"{args.file}: " + ("all checks passed" if not problems else f"{len(problems)} difference(s)"))
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
